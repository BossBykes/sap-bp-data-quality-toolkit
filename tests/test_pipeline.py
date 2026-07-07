from pathlib import Path

from openpyxl import load_workbook
import pandas as pd
import yaml

from sap_bp_dq.pipeline import run_pipeline
from sap_bp_dq.report import render_report


EXPECTED_WORKBOOK_SHEETS = {
    "Summary",
    "Cleaned Records",
    "Issues",
    "Exact Duplicates",
    "Fuzzy Duplicates",
}


def _write_config(path: Path, *, fuzzy_enabled: bool = True) -> None:
    config = {
        "required_fields": ["bp_id", "bp_type", "name", "country", "city"],
        "allowed_bp_types": ["PERSON", "COMPANY"],
        "country_rules": {"allowed": ["DE", "US"], "uppercase": True},
        "email_rules": {"enabled": True},
        "phone_rules": {"enabled": True, "min_digits": 7},
        "dedup_rules": {
            "exact_keys": ["name", "city", "country"],
            "fuzzy_enabled": fuzzy_enabled,
            "fuzzy_threshold": 90,
            "fuzzy_keys": ["name", "city"],
            "fuzzy_allow_cross_country": False,
        },
    }
    path.write_text(yaml.safe_dump(config), encoding="utf-8")


def test_report_renders_empty_sections(tmp_path):
    out_path = tmp_path / "report.html"

    render_report(
        out_path=out_path,
        total_rows=0,
        issues=pd.DataFrame(columns=["row_index", "bp_id", "field", "issue", "severity"]),
        exact_dups=pd.DataFrame(),
        fuzzy_pairs=pd.DataFrame(),
    )

    html = out_path.read_text(encoding="utf-8")
    assert "No validation issues found." in html
    assert "No exact duplicate rows found." in html
    assert "No fuzzy duplicate pairs found." in html


def test_pipeline_runs_with_fuzzy_disabled(tmp_path):
    input_path = tmp_path / "business_partners.csv"
    config_path = tmp_path / "config.yaml"
    out_dir = tmp_path / "output"

    pd.DataFrame(
        [
            {
                "bp_id": "BP1",
                "bp_type": "COMPANY",
                "name": "RWE AG",
                "email": "",
                "phone": "0201123456",
                "country": "de",
                "city": "Essen",
            },
            {
                "bp_id": "BP2",
                "bp_type": "COMPANY",
                "name": "RWE AG",
                "email": "",
                "phone": "0201123457",
                "country": "DE",
                "city": "Essen",
            },
        ]
    ).to_csv(input_path, index=False)
    _write_config(config_path, fuzzy_enabled=False)

    results = run_pipeline(input_path=input_path, config_path=config_path, out_dir=out_dir)

    assert Path(results["cleaned_csv"]).exists()
    assert Path(results["issues_csv"]).exists()
    assert Path(results["report_html"]).exists()
    assert Path(results["excel_report"]).exists()
    assert Path(results["log_file"]).exists()

    report_html = Path(results["report_html"]).read_text(encoding="utf-8")
    assert "No fuzzy duplicate pairs found." in report_html

    workbook = load_workbook(results["excel_report"], read_only=True)
    assert set(workbook.sheetnames) == EXPECTED_WORKBOOK_SHEETS


def test_small_full_pipeline_run(tmp_path):
    input_path = tmp_path / "business_partners.csv"
    config_path = tmp_path / "config.yaml"
    out_dir = tmp_path / "output"

    pd.DataFrame(
        [
            {
                "bp_id": "BP1",
                "bp_type": "PERSON",
                "name": None,
                "email": "bad-email",
                "phone": "12",
                "country": "zz",
                "city": "",
            },
            {
                "bp_id": "BP2",
                "bp_type": "PERSON",
                "name": "Anna Schmidt",
                "email": "anna@example.com",
                "phone": "0201123456",
                "country": "DE",
                "city": "Essen",
            },
        ]
    ).to_csv(input_path, index=False)
    _write_config(config_path)

    results = run_pipeline(input_path=input_path, config_path=config_path, out_dir=out_dir)

    issues = pd.read_csv(results["issues_csv"])
    assert {"missing_required", "invalid_email_format", "phone_too_short", "invalid_country"}.issubset(
        set(issues["issue"])
    )
    assert Path(results["report_html"]).exists()
    assert Path(results["excel_report"]).exists()


def test_pipeline_creates_excel_workbook_with_expected_sheets(tmp_path):
    input_path = tmp_path / "business_partners.csv"
    config_path = tmp_path / "config.yaml"
    out_dir = tmp_path / "output"

    pd.DataFrame(
        [
            {
                "bp_id": "BP1",
                "bp_type": "COMPANY",
                "name": "RWE AG",
                "email": "",
                "phone": "0201123456",
                "country": "DE",
                "city": "Essen",
            },
            {
                "bp_id": "BP2",
                "bp_type": "COMPANY",
                "name": "RWE AG",
                "email": "",
                "phone": "0201123457",
                "country": "DE",
                "city": "Essen",
            },
        ]
    ).to_csv(input_path, index=False)
    _write_config(config_path)

    results = run_pipeline(input_path=input_path, config_path=config_path, out_dir=out_dir)

    workbook = load_workbook(results["excel_report"], read_only=True)
    assert set(workbook.sheetnames) == EXPECTED_WORKBOOK_SHEETS

    summary = workbook["Summary"]
    metrics = {row[0].value: row[1].value for row in summary.iter_rows(min_row=2)}
    assert metrics["total_input_rows"] == 2
    assert metrics["exact_duplicate_rows"] == 2
    assert "generated_at_utc" in metrics
